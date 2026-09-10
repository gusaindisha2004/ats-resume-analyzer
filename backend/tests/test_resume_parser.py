"""File validation and text extraction.

File type is decided by content signature, never by the filename — a renamed
file must not get through.
"""

import io

import pytest

from backend.core.config import MAX_FILE_SIZE_BYTES
from backend.services.resume_parser import (
    FileParsingError,
    FileValidationError,
    detect_file_type,
    extract_text_from_docx,
    parse_resume_file,
    validate_file,
)

OLE2_HEADER = bytes.fromhex('d0cf11e0a1b11ae1')
ZIP_HEADER = bytes.fromhex('504b0304')


def make_docx(paragraphs=('Hello world',), table_rows=None) -> bytes:
    from docx import Document

    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDetectFileType:
    def test_detects_pdf_by_signature(self):
        assert detect_file_type(b'%PDF-1.7\nrest of file') == 'pdf'

    def test_detects_docx_by_its_word_part(self):
        assert detect_file_type(make_docx()) == 'docx'

    def test_detects_legacy_doc_container(self):
        assert detect_file_type(OLE2_HEADER + b'padding') == 'doc'

    def test_plain_text_is_not_a_document(self):
        assert detect_file_type(b'Alex Rivera, Backend Engineer') is None

    def test_a_zip_without_a_word_part_is_rejected(self):
        """.xlsx and .pptx are ZIPs too — the container alone proves nothing."""
        buffer = io.BytesIO()
        import zipfile

        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('xl/workbook.xml', '<workbook/>')
        assert detect_file_type(buffer.getvalue()) is None

    def test_truncated_zip_does_not_raise(self):
        assert detect_file_type(ZIP_HEADER + b'\x00' * 10) is None

    def test_executable_is_rejected(self):
        assert detect_file_type(b'MZ\x90\x00\x03' + b'\x00' * 60) is None

    def test_empty_input_is_rejected(self):
        assert detect_file_type(b'') is None


class TestValidateFile:
    def test_accepts_a_docx(self):
        is_valid, message, file_type = validate_file(make_docx(), 'resume.docx')
        assert is_valid and file_type == 'docx' and message == ''

    def test_accepts_a_pdf(self):
        is_valid, _, file_type = validate_file(b'%PDF-1.4 content', 'resume.pdf')
        assert is_valid and file_type == 'pdf'

    def test_rejects_an_empty_file(self):
        is_valid, message, file_type = validate_file(b'', 'resume.pdf')
        assert not is_valid and file_type is None
        assert 'empty' in message.lower()

    def test_rejects_an_oversized_file(self):
        oversized = b'%PDF-' + b'x' * MAX_FILE_SIZE_BYTES
        is_valid, message, _ = validate_file(oversized, 'big.pdf')
        assert not is_valid
        assert 'exceeds' in message.lower()

    def test_rejects_legacy_doc_with_conversion_advice(self):
        is_valid, message, _ = validate_file(OLE2_HEADER + b'x', 'old.doc')
        assert not is_valid
        assert '.docx' in message

    def test_a_renamed_file_is_judged_by_content(self):
        """Extension says .pdf; content says otherwise."""
        is_valid, _, _ = validate_file(b'not a pdf at all', 'resume.pdf')
        assert not is_valid

    def test_always_returns_a_three_tuple(self):
        """Regression guard: the empty-file branch once returned two values,
        which raised a ValueError in the caller."""
        for data in (b'', b'%PDF-1.4', b'garbage', OLE2_HEADER, make_docx()):
            assert len(validate_file(data, 'f')) == 3


class TestDocxExtraction:
    def test_extracts_paragraph_text(self):
        text = extract_text_from_docx(make_docx(['ALEX RIVERA', 'Backend Engineer']))
        assert 'ALEX RIVERA' in text
        assert 'Backend Engineer' in text

    def test_extracts_table_cell_text(self):
        """Resumes built on table layouts are common — their text still counts."""
        text = extract_text_from_docx(
            make_docx(['Header'], table_rows=[['Python', 'FastAPI']])
        )
        assert 'Python' in text and 'FastAPI' in text

    def test_empty_document_raises(self):
        with pytest.raises(FileParsingError):
            extract_text_from_docx(make_docx([]))

    def test_corrupt_input_raises_file_parsing_error(self):
        with pytest.raises(FileParsingError):
            extract_text_from_docx(b'not a docx')


class TestParseResumeFile:
    def test_returns_text_and_metadata(self):
        data = make_docx(['ALEX RIVERA', 'Developed FastAPI services'])
        text, metadata = parse_resume_file(data, 'resume.docx')

        assert 'ALEX RIVERA' in text
        assert metadata['file_type'] == 'docx'
        assert metadata['filename'] == 'resume.docx'
        assert metadata['file_size_bytes'] == len(data)
        assert metadata['text_length'] == len(text)
        assert metadata['success'] is True

    def test_invalid_file_raises_validation_error(self):
        with pytest.raises(FileValidationError):
            parse_resume_file(b'plain text', 'resume.txt')

    def test_empty_file_raises_validation_error_not_value_error(self):
        with pytest.raises(FileValidationError):
            parse_resume_file(b'', 'resume.docx')

    def test_legacy_doc_raises_validation_error(self):
        with pytest.raises(FileValidationError):
            parse_resume_file(OLE2_HEADER + b'x', 'resume.doc')
